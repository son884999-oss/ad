"""Build editable PPTX and an independently rendered PDF from shared content.

Run from any directory. Dependencies are isolated in the workspace .pt_deps.
Rebuilds this folder's presentation, PDF, previews and check records.
"""
from pathlib import Path
import sys
import json
import re
import hashlib

OUT = Path(__file__).resolve().parent
ROOT = OUT.parent
sys.path.insert(0, str(ROOT / '.pt_deps'))
import pymupdf as fitz
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from PIL import Image

INK, TEAL, MINT = '102D34', '087E80', 'BDF4D5'
PAPER, WHITE, CORAL = 'F3F6F2', 'FFFFFF', 'FF936F'
MUTED, LINE, FAINT = '526B70', 'D9E3DF', 'B5CFCC'
MSS = 'https://www.mss.go.kr/site/smba/ex/bbs/View.do?bcIdx=1066231&cbIdx=86'
NABO = 'https://www.nabo.go.kr/board/file/down.do?fid=33318912'
COMMIT = '75669cb353ebf235acf938b693afd29da6e842b2'
APP_DOC = f'https://github.com/son884999-oss/ad/blob/{COMMIT}/ad_prototype/docs/README.md'
CNI = ('https://www.cni.re.kr/main/contents/report/down?addType=2&docId=1024E10035'
       '&gcd=BB0000002137&projectName=%EC%B6%A9%EB%82%A8%EB%A6%AC%ED%8F%AC%ED%8A%B8-%28%EC%A0%9C397%ED%98%B8%29'
       '+%EC%B6%A9%EC%B2%AD%EB%82%A8%EB%8F%84+%EC%86%8C%EC%83%81%EA%B3%B5%EC%9D%B8'
       '+%EB%94%94%EC%A7%80%ED%84%B8%EC%A0%84%ED%99%98+%EC%B6%94%EC%A7%84%EB%B0%A9%EC%95%88'
       '&projectType=15&seq=1')
FONT = Path('C:/Windows/Fonts/malgun.ttf')
BOLD = Path('C:/Windows/Fonts/malgunbd.ttf')
assert FONT.exists() and BOLD.exists(), 'Windows Malgun Gothic fonts required'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(16), Inches(9)
prs.core_properties.title = '홍보 잇다 | 기획 의도'
prs.core_properties.subject = '소상공인의 홍보 콘텐츠 제작 부담 완화'
prs.core_properties.author = '홍보 잇다 프로젝트'
pdf = fitz.open()
records = []
evidence_records = []


def rgb(c):
    return tuple(int(c[i:i+2], 16) / 255 for i in (0, 2, 4))


def shape(s, p, x, y, w, h, color, kind='rect'):
    typ = {'rect': MSO_SHAPE.RECTANGLE, 'round': MSO_SHAPE.ROUNDED_RECTANGLE,
           'circle': MSO_SHAPE.OVAL}[kind]
    sh = s.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color)
    sh.line.fill.background()
    if kind == 'round':
        sh.adjustments[0] = 0.07
    box = fitz.Rect(x * 72, y * 72, (x + w) * 72, (y + h) * 72)
    if kind == 'circle':
        p.draw_oval(box, color=None, fill=rgb(color))
    else:
        p.draw_rect(box, color=None, fill=rgb(color), radius=0.07 if kind == 'round' else None)
    return sh


def txt(s, p, value, x, y, w, h, size=24, color=INK, bold=False, url=None):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = line
        para.font.name = '맑은 고딕'
        para.font.size = Pt(size)
        para.font.bold = bold
        para.font.color.rgb = RGBColor.from_string(color)
        para.space_before = para.space_after = Pt(0)
        para.line_spacing = 1.16
        if url:
            for run in para.runs:
                run.hyperlink.address = url
    rect = fitz.Rect(x * 72, y * 72, (x + w) * 72, (y + h) * 72)
    spare = p.insert_textbox(rect, value, fontsize=size, fontname='b' if bold else 'r',
                             lineheight=1.16, color=rgb(color))
    if spare < 0:
        raise ValueError(f'Text overflow ({spare:.1f}): {value}')
    if url:
        p.insert_link({'kind': fitz.LINK_URI, 'from': rect, 'uri': url})
    records.append({'slide': len(prs.slides), 'text': value, 'size': size,
                    'bounds': [x, y, w, h], 'url': url})
    return box


def page(n, label, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg, fg = (INK, FAINT) if dark else (PAPER, MUTED)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
    p = pdf.new_page(width=1152, height=648)
    p.draw_rect(p.rect, fill=rgb(bg), color=None)
    p.insert_font(fontname='r', fontfile=str(FONT))
    p.insert_font(fontname='b', fontfile=str(BOLD))
    txt(s, p, '홍보 잇다', .65, .36, 3, .44, 18, MINT if dark else TEAL, True)
    txt(s, p, label, 8.2, .41, 7.1, .3, 11.5, fg)
    shape(s, p, .65, 8.46, 14.7, .012, '36565D' if dark else LINE)
    txt(s, p, 'PLANNING INTENT  /  2026.10.24', .65, 8.58, 9, .24, 10, fg)
    txt(s, p, f'{n:02d} / 08', 14.35, 8.57, 1, .27, 11, fg)
    return s, p


def note(s, text):
    s.notes_slide.notes_text_frame.text = text


def evidence_image(s, p, name, source, index, bounds, x, y, w, h):
    """Render an unaltered rectangular crop from a retained official PDF."""
    folder = OUT / 'evidence'
    source_path = folder / source
    doc = fitz.open(source_path)
    region = fitz.Rect(bounds)
    value = doc[index].get_text(clip=region)
    crop_path = folder / name
    doc[index].get_pixmap(matrix=fitz.Matrix(4, 4), clip=region).save(crop_path)
    iw, ih = region.width, region.height
    scale = min(w/iw, h/ih)
    ww, hh = iw*scale, ih*scale
    xx, yy = x+(w-ww)/2, y+(h-hh)/2
    s.shapes.add_picture(str(crop_path), Inches(xx), Inches(yy), width=Inches(ww), height=Inches(hh))
    p.insert_image(fitz.Rect(xx*72, yy*72, (xx+ww)*72, (yy+hh)*72), filename=str(crop_path))
    evidence_records.append({'slide':len(prs.slides), 'image':name, 'source':source,
                             'pdf_page':index+1, 'crop_points':bounds,
                             'source_text':value, 'image_sha256':hashlib.sha256(crop_path.read_bytes()).hexdigest()})
    doc.close()


# 1 — Large editorial cover with a native vector connection motif.
s, p = page(1, '01  /  기획 의도', True)
shape(s, p, .7, 1.53, 2.16, .4, TEAL, 'round')
txt(s, p, 'SMALL BUSINESS × AI', .84, 1.59, 2, .25, 10, WHITE, True)
txt(s, p, '가게의 가치를,\n고객에게 잇다.', .68, 2.26, 10.3, 2.42, 56, WHITE, True)
txt(s, p, '제품 사진에서 시작하는\n소상공인 홍보 콘텐츠 제작', .75, 5.18, 9.3, 1.3, 27, MINT)
txt(s, p, '홍보 도구 사용에 부담을 느끼는 소상공인을 위한 기획', .75, 7.31, 12, .48, 16, FAINT)
shape(s, p, 11.18, 2.15, 3.65, 3.65, TEAL, 'circle')
shape(s, p, 10.85, 5.61, 1.04, 1.04, CORAL, 'circle')
for y, label, bg, fg in [(2.69, '가게', MINT, INK), (5.22, '고객', WHITE, INK)]:
    shape(s, p, 11.86, y, 2.35, 1.06, bg, 'round')
    txt(s, p, label, 12.42, y+.22, 1.45, .65, 27, fg, True)
shape(s, p, 12.98, 3.75, .08, 1.47, MINT)
shape(s, p, 12.77, 4.31, .5, .5, CORAL, 'circle')
note(s, '약 15초. 홍보 잇다는 홍보 도구 사용에 부담을 느끼는 소상공인을 위해 기획했습니다. '
     '제품 사진과 간단한 정보에서 시작해 홍보 콘텐츠를 만들고, 가게가 고객과 연결될 기회를 넓히는 것이 목적입니다. '
     '표지와 슬로건은 팀의 기획 문안이며 확정된 메인 로고가 아닙니다.')

# 2 — Independent statistic cards; no misleading shared chart or funnel.
s, p = page(2, '02  /  문제와 근거')
txt(s, p, '기술을 제공하는 것에서,\n활용을 돕는 것으로.', .7, 1.15, 14.6, 1.46, 36, INK, True)
for x, bg in [(.7, WHITE), (8.1, INK)]:
    shape(s, p, x, 2.96, 7.2, 3.38, bg, 'round')
txt(s, p, '디지털·스마트 기술 활용 비중', 1.02, 3.2, 6.55, .55, 22, TEAL, True)
txt(s, p, '27.2%', 1.02, 3.91, 6.2, 1.4, 66, INK, True)
txt(s, p, '2024년 기준 소상공인실태조사', 1.02, 5.45, 6.55, .65, 22, MUTED)
txt(s, p, '기술 미도입 응답자 · 이유 1순위', 8.43, 3.2, 6.5, .55, 22, MINT, True)
txt(s, p, '45.6%', 8.43, 3.91, 6.2, 1.4, 66, WHITE, True)
txt(s, p, '인력·지식 부재', 8.43, 5.45, 6.5, .65, 24, WHITE, True)
txt(s, p, '기획 판단  /  도구를 활용하는 과정의 부담에 주목했습니다.', .8, 6.63, 14.4, .58, 23, INK, True)
txt(s, p, '서로 다른 조사입니다. 홍보 도구 이용률·컴퓨터 능력·서비스 수요를 뜻하지 않습니다. 원문: 6~7쪽', .8, 7.39, 14.4, .35, 13, MUTED)
txt(s, p, '[S1] 중기부, 2024년 기준 실태조사 결과 발표, 2026.03.13.', .8, 7.93, 7.15, .3, 11, MUTED, url=MSS)
txt(s, p, '[S2] 국회예산정책처, 2025.11., 97쪽 각주 65 · 소진공 조사 재인용', 8.1, 7.93, 7.2, .3, 11, MUTED, url=NABO)
note(s, '약 35초. 중기부 조사에서 2024년 기준 디지털·스마트 기술 활용 비중은 27.2%였습니다. '
     '별도의 디지털 전환 조사에서는 기술 미도입 응답자의 미도입 이유 1순위로 인력과 지식 부족이 45.6%를 차지했습니다. '
     '두 수치는 조사대상이 다릅니다. 우리는 이를 개발 배경으로 삼아, 도구를 제공하는 것과 함께 활용 과정의 부담을 낮추는 방향을 선택했습니다.\n'
     'S1: 11개 산업 약 4만개 업체 대상 조사. 약 4만개는 정확한 유효응답 수 아님. '+MSS+'\n'
     'S2: 국회예산정책처 본문 97쪽 각주 65, PDF 120쪽. 소진공 2024년 소상공인 디지털 전환 실태조사(2024.12.) 재인용. '
     '최초 조사방법 전체 미확인. 전체 소상공인 비율 아님. '+NABO)

# 3 — A single simple journey with strong hierarchy.
s, p = page(3, '03  /  서비스의 목표 흐름')
txt(s, p, '사진과 정보에서,\n홍보물 완성까지.', .7, 1.17, 14.5, 1.46, 38, INK, True)
items = [(0.7, '01', '사진·정보 입력', '사진 1장·제품 정보\n자동 채우기 지원'),
         (5.75, '02', '홍보물 제작', '포스터·10초 영상\n홍보 글'),
         (10.8, '03', '확인·활용', '결과 확인·복사\n저장·공유')]
for x, num, title, detail in items:
    shape(s, p, x, 3.08, 4.5, 3.25, INK if num=='02' else WHITE, 'round')
    fg = WHITE if num=='02' else INK
    txt(s, p, num, x+.3, 3.35, 1, .4, 17, MINT if num=='02' else TEAL, True)
    txt(s, p, title, x+.3, 4.03, 3.95, .67, 32, fg, True)
    txt(s, p, detail, x+.3, 5.04, 3.95, 1.05, 23, FAINT if num=='02' else MUTED)
for x in [5.28, 10.33]:
    txt(s, p, '→', x, 4.37, .45, .6, 24, TEAL, True)
txt(s, p, 'AI는 사진 분석·글 작성·선택적 배경 변경, 포스터·영상은 브라우저 제작', .8, 6.69, 14.4, .6, 22, INK, True)
txt(s, p, '팀 문서 기준: 75669cb · 사용 부담을 줄이는 통합 흐름 / 참고 연구 원문: 8쪽', .8, 7.41, 14.4, .34, 13, MUTED, url=APP_DOC)
txt(s, p, '[S3] 참고: 충남연구원, 「충청남도 소상공인 디지털전환 추진방안」, 2024.09.06., 4~5쪽 · 단계별 지원과 소비자 연결', .8, 7.95, 14.5, .28, 11, MUTED, url=CNI)
note(s, '약 35초. 사용자는 제품 사진을 올리고 필요한 정보를 입력합니다. 자동 채우기를 누르면 AI가 사진을 분석해 입력을 돕습니다. '
     'AI는 홍보 글과 선택적인 배경 변경에 활용하고, 포스터와 10초 영상은 브라우저에서 제작합니다. 결과를 확인하고 복사·저장·공유하는 흐름입니다. '
     '사용자 수준에 맞춘 지원과 소비자 연결을 강조한 충남연구원 보고서도 참고했습니다. '
     '여러 제작 과정을 간단한 흐름으로 이어주는 것이 기획의 핵심입니다.\n'
     '구현 설명은 공유된 커밋 75669cb의 팀 문서 기준이며 이번 작업에서 앱을 실행해 검증한 것은 아닙니다. '
     '현재 기본 영상을 외부 AI 영상 생성 API의 출력으로 설명하지 않습니다. '+APP_DOC+'\n'
     'S3 저자 이민정·전춘복, 충남리포트 제397호, PDF 4~5쪽. 지역 정책연구를 참고한 팀의 해석으로서 앱 효과 검증 아님. '+CNI)

# 4 — Principles and a concrete validation strip.
s, p = page(4, '04  /  설계 원칙과 검증', True)
txt(s, p, '홍보를 시작할 수 있도록.', .7, 1.22, 14.6, .95, 42, WHITE, True)
txt(s, p, '제작 부담 완화와 홍보 기회 확대를 목표로 합니다.', .75, 2.48, 14.5, .66, 25, MINT)
for x, num, title, detail in [(.7, '01', '쉬운 화면', '큰 글씨·쉬운 말\n말로 입력·듣기'),
                             (5.75, '02', '제품에 충실한 표현', '형태·색상·라벨 보존\n확인되지 않은 특성 배제'),
                             (10.8, '03', '활용 가능한 결과', '결과 확인에서\n저장까지 연결')]:
    shape(s, p, x, 3.55, 4.5, 2.62, '1C4148', 'round')
    txt(s, p, num, x+.27, 3.8, .6, .35, 13, CORAL, True)
    txt(s, p, title, x+.27, 4.37, 4, .59, 23, WHITE, True)
    txt(s, p, detail, x+.27, 5.13, 4, .88, 22, FAINT)
shape(s, p, .7, 6.63, 14.6, 1.17, MINT, 'round')
txt(s, p, '검증 계획', .99, 6.92, 2, .44, 20, TEAL, True)
txt(s, p, '완성·저장 여부  /  소요시간  /  도움 요청  /  재생성 이유', 3.16, 6.96, 11.65, .52, 22, INK, True)
txt(s, p, '팀의 설계 목표와 검증 계획입니다. 시간 절감·매출 증가 등의 효과는 아직 검증하지 않았습니다.', .8, 8.03, 14.4, .28, 11, FAINT)
note(s, '약 35초. 설계 기준은 쉬운 화면, 제품에 충실한 표현, 활용 가능한 결과물입니다. '
     '팀 문서에는 큰 글씨와 쉬운 말, 음성 입력과 듣기를 제공하는 쉬운 화면이 기록되어 있습니다. '
     '기대하는 변화는 제작 부담 완화와 홍보 기회 확대입니다. 실제 효과는 사용자가 홍보물을 완성하고 저장하는지, '
     '시간이 얼마나 걸리고 어떤 도움이 필요한지 관찰해 확인하겠습니다. '
     '팀의 설계 목표이며 해당 보존 규칙이나 결과 품질이 전체 서비스에서 보장된다는 뜻은 아닙니다. '
     '실제 사용자 테스트 결과·정량 개선율을 주장하지 않습니다. 구현 설명 기준: '+APP_DOC)

# 5 — Source appendix. Entire source rows link to primary documents.
s, p = page(5, '05  /  출처 부록 · 확인일 2026.10.24.')
txt(s, p, '근거와 기획을 구분했습니다.', .7, 1.21, 14.6, .9, 36, INK, True)
refs = [
    ('S1', '디지털·스마트 기술 활용 27.2%',
     '중소벤처기업부 · 2024년 기준 소상공인실태조사 결과 발표\n2026.03.13. 발표 / 11개 산업 약 4만개 업체 대상 / 공식 발표 직접 확인', MSS),
    ('S2', '미도입 이유 1순위 · 인력·지식 부재 45.6%',
     '국회예산정책처 · 소상공인 지원 재정사업 평가 / 2025.11.\n본문 97쪽 각주 65 (PDF 120쪽) / 소진공 2024년 조사 재인용 / 기술 미도입 응답자', NABO),
    ('S3', '단계별 지원과 소비자 연결',
     '이민정·전춘복 · 충남연구원 · 충청남도 소상공인 디지털전환 추진방안\n충남리포트 제397호 / 2024.09.06. / 4~5쪽 / 지역 정책연구의 제안을 기획에 참고', CNI),
]
for y, (num, title, detail, url) in zip([2.59, 4.13, 5.67], refs):
    shape(s, p, .7, y, 14.6, 1.36, WHITE, 'round')
    shape(s, p, .94, y+.23, .69, .69, TEAL, 'circle')
    txt(s, p, num, 1.08, y+.39, .5, .3, 12, WHITE, True)
    txt(s, p, title, 1.94, y+.13, 12.95, .43, 19, INK, True, url)
    txt(s, p, detail, 1.94, y+.68, 12.95, .56, 12.2, MUTED, url=url)
txt(s, p, '자료명 클릭 → 원문  /  상세 출처·확인 범위: sources.md', .8, 7.26, 14.4, .32, 12, TEAL, True)
txt(s, p, '[S4] 구현 설명 기준: son884999-oss/ad · 커밋 75669cb의 팀 문서 (클릭)', .8, 7.65, 14.4, .31, 11.5, TEAL, url=APP_DOC)
txt(s, p, '최초 조사방법·원응답 데이터와 앱 실행은 미검증. 통계 근거·팀 문서·기획 해석을 구분했습니다.', .8, 8.08, 14.4, .29, 11, MUTED)
note(s, '질의응답용. S1은 공식 발표 직접 확인, S2는 국회예산정책처 보고서를 통한 소진공 조사 재인용. '
     'S3는 충남 지역 정책연구의 제안을 기획에 참고. 통계끼리 결합하지 않았으며 팀의 기획 문안과 구분했습니다.\n'
     'S1 '+MSS+'\nS2 '+NABO+'\nS3 '+CNI+'\nS4 구현 설명 기준 '+APP_DOC)

# 6–8 — Actual source evidence, with unmodified cropped pixels and context.
s, p = page(6, '06  /  근거 원문 S1 · 중소벤처기업부')
txt(s, p, '27.2%의 근거를 직접 확인합니다.', .7, 1.19, 14.6, .86, 35, INK, True)
txt(s, p, '2024년 기준 소상공인실태조사 결과 발표 · 공식 첨부 PDF 2쪽', .76, 2.31, 14.4, .52, 21, TEAL, True)
shape(s, p, .7, 3.12, 14.6, 2.76, WHITE, 'round')
txt(s, p, '원문 문단 발췌 · 원본 비율 유지', 1.02, 3.32, 13.8, .32, 12, MUTED)
evidence_image(s, p, 's1_original.png', 'mss_2024_release.pdf', 1,
               [55, 148, 545, 224], 1.02, 3.92, 13.92, 1.72)
txt(s, p, '읽는 기준', .8, 6.22, 2.2, .44, 21, TEAL, True)
txt(s, p, '2024년 기준의 기술 활용 현황입니다.\n홍보 도구 이용률이나 사용능력을 측정한 지표는 아닙니다.', 3.03, 6.2, 12.05, 1.1, 23, INK, True)
txt(s, p, '문단 안의 증감·활용 유형 수치는 원문 보존을 위해 함께 보이며, 이번 발표의 추가 주장으로 사용하지 않습니다.', .8, 7.51, 14.4, .34, 12, MUTED)
txt(s, p, '[S1] 중기부, 2026.03.13. 발표 / 공식 첨부 PDF 2쪽 · 클릭하여 공식 게시물 열기', .8, 8.04, 14.4, .27, 11, MUTED, url=MSS)
note(s, '근거 확인용 부록. 중기부 공식 게시물 첨부 보도자료 PDF 2쪽의 해당 문단을 이미지로 직접 발췌했다. '
     '원문은 다시 쓰거나 수치·문구를 수정하지 않았다. 문단 내 다른 유형 비율과 증감은 맥락 보존용이며 '
     '본 발표는 27.2%만 사용한다. 다운로드 URL과 파일 해시·좌표는 evidence 폴더 기록 참조. '+MSS)

s, p = page(7, '07  /  근거 원문 S2 · 국회예산정책처')
txt(s, p, '45.6%는 누구의 응답일까요?', .7, 1.19, 14.6, .86, 35, INK, True)
txt(s, p, '소상공인 지원 재정사업 평가 · 본문 97쪽 각주 65 / PDF 120쪽', .76, 2.31, 14.4, .52, 21, TEAL, True)
shape(s, p, .7, 3.12, 14.6, 2.32, WHITE, 'round')
txt(s, p, '각주 앞부분 2줄 발췌 · 이후 내용 생략 · 원본 비율 유지', 1.02, 3.32, 13.8, .32, 12, MUTED)
evidence_image(s, p, 's2_original.png', 'nabo_2025.pdf', 119,
               [80, 597, 462, 622], 1.02, 3.98, 13.92, 1.06)
shape(s, p, .7, 5.8, 14.6, 1.58, INK, 'round')
txt(s, p, '응답집단', 1.02, 6.03, 2.2, .43, 19, MINT, True)
txt(s, p, '디지털 기술(기기) 미도입 응답자', 3.38, 6.01, 11.3, .54, 26, WHITE, True)
txt(s, p, '질문 조건', 1.02, 6.77, 2.2, .43, 19, MINT, True)
txt(s, p, '미도입 이유 1순위', 3.38, 6.75, 11.3, .48, 23, WHITE)
txt(s, p, '전체 소상공인 비율이 아닙니다. 소진공의 2024년 조사를 재인용한 자료이며, 최초 조사방법 전체는 미확인입니다.', .8, 7.68, 14.4, .34, 12, MUTED)
txt(s, p, '[S2] 국회예산정책처, 2025.11., 본문 97쪽 각주 65 / PDF 120쪽 · 클릭하여 원문 열기', .8, 8.12, 14.4, .27, 11, MUTED, url=NABO)
note(s, '근거 확인용 부록. 공식 보고서 PDF 120쪽 각주 65의 첫 2줄만 원본 이미지로 발췌했다. '
     '자금 부족 등의 뒤 내용과 최초 조사 출처 문장은 발췌 영역 밖이며 생략 사실을 표시했다. '
     '최초 자료: 소상공인시장진흥공단, 2024년 소상공인 디지털 전환 실태조사, 2024.12. '
     '전체 소상공인 비율이나 컴퓨터 사용능력 지표로 해석하지 않는다. '+NABO)

s, p = page(8, '08  /  근거 원문 S3 · 충남연구원')
txt(s, p, '수준을 진단하고, 단계에 맞춰 돕는다.', .7, 1.19, 14.6, .86, 34, INK, True)
txt(s, p, '충청남도 소상공인 디지털전환 추진방안 · 충남리포트 제397호', .76, 2.31, 14.4, .52, 21, TEAL, True)
shape(s, p, .7, 3.08, 14.6, 1.76, WHITE, 'round')
txt(s, p, '5쪽 정책 제언 · 원문 2곳 발췌 / 두 문장 사이의 설명은 생략', 1.02, 3.24, 13.8, .32, 12, MUTED)
evidence_image(s, p, 's3_level_original.png', 'cni_397.pdf', 4,
               [62, 171, 438, 188], 1.02, 3.63, 13.92, .38)
evidence_image(s, p, 's3_fit_original.png', 'cni_397.pdf', 4,
               [85, 207, 545, 223], 1.02, 4.18, 13.92, .38)
txt(s, p, '단계별 지원 제안의 요지 · 아래는 원문을 요약해 정리한 내용입니다.', .8, 5.06, 14.4, .36, 13, TEAL, True)
for y, label, detail in [
    (5.58, '도입 준비가 필요한 단계', '기초 교육·이용 상담·현장 보조로 시작을 지원'),
    (6.16, '기술은 알지만 활용이 낮은 단계', '활용 지도·결제 환경·온라인 판매·배송 기반을 지원'),
    (6.74, '스스로 활용·확장 가능한 단계', '업종별 장비·물류 체계·데이터 활용으로 확장을 지원'),
]:
    shape(s, p, .7, y, 14.6, .51, WHITE, 'round')
    txt(s, p, label, .97, y+.08, 5.1, .36, 17, INK, True)
    txt(s, p, detail, 6.25, y+.08, 8.7, .36, 17, MUTED)
txt(s, p, '기획에 적용  /  큰 글씨·쉬운 말·입력 도움으로, 익숙하지 않은 사용자의 시작을 돕습니다.', .8, 7.57, 14.4, .4, 16, INK, True)
txt(s, p, '[S3] 이민정·전춘복, 충남연구원, 2024.09.06., 5쪽 · 지역 정책 제안이며 앱 효과 검증은 아님', .8, 8.13, 14.4, .27, 11, MUTED, url=CNI)
note(s, '근거 확인용 부록. 공식 PDF 5쪽에서 수준 파악의 필요성과 체크리스트 점수에 따른 적합사업 선택에 관한 '
     '두 문장을 별도 원본 이미지로 발췌하고 사이 설명의 생략 사실을 표시했다. '
     '그 아래 단계별 정의와 지원 제안을 세 행으로 요약했다. 기초 단계에는 교육·상담·현장 보조, '
     '인식은 있으나 활용이 낮은 단계에는 활용 지도·결제 환경·온라인 판매배송, '
     '역량을 보유한 단계에는 업종별 장비·물류·데이터 활용을 제안한다는 맥락이다. '
     '행 이름과 지원 내용은 편집자의 요약이며 원문 그대로의 표가 아니다. '
     '지역 조사 추정 비율은 전국 수치와 혼동하지 않도록 추가하지 않았다. '
     '쉬운 화면에 대한 연결은 팀의 기획 해석이며 연구원이 우리 앱을 검증한 것은 아니다. '+CNI)

# Save and render independently; semantic checks compare PPTX/PDF contents.
target = OUT / 'planning_intent.pptx'
prs.save(target)
pdf.subset_fonts()
pdf.save(OUT / 'planning_intent.pdf', garbage=4, deflate=True)
pdf.close()
doc = fitz.open(OUT / 'planning_intent.pdf')
loaded = Presentation(target)
assert len(doc) == len(loaded.slides) == 8
norm = lambda text: re.sub(r'\s+', '', text)
for i, (ps, pd) in enumerate(zip(loaded.slides, doc)):
    ptext = norm(pd.get_text())
    for sh in ps.shapes:
        if sh.has_text_frame and sh.text.strip():
            assert norm(sh.text) in ptext, f'PDF text mismatch slide {i+1}: {sh.text}'
    assert ps.has_notes_slide
for rec in records:
    x, y, w, h = rec['bounds']
    assert x >= 0 and y >= 0 and x+w <= 16.001 and y+h <= 9.001
assert '27.2%' in doc[1].get_text() and '45.6%' in doc[1].get_text()
assert '미도입 응답자' in doc[1].get_text()
assert all(list(doc[i].get_links()) for i in [1, 2, 4, 5, 6, 7])
assert len(evidence_records) == 4
with __import__('zipfile').ZipFile(target) as archive:
    hashes = {hashlib.sha256(archive.read(n)).hexdigest() for n in archive.namelist() if n.startswith('ppt/media/')}
    assert all(item['image_sha256'] in hashes for item in evidence_records)
(OUT/'evidence/crop_manifest.json').write_text(json.dumps(evidence_records, ensure_ascii=False, indent=2), encoding='utf-8')

sheet = Image.new('RGB', (1600, 1800), '#D9E3DF')
for i, p in enumerate(doc):
    pix = p.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
    im = Image.frombytes('RGB', [pix.width, pix.height], pix.samples)
    im.save(OUT / f'slide_{i+1:02d}.png')
    thumb = im.resize((800, 450), Image.Resampling.LANCZOS)
    sheet.paste(thumb, ((i % 2)*800, (i // 2)*450))
sheet.save(OUT / 'preview.png')
report = {'slides': 8, 'pptx_pdf_text_match': True, 'text_boxes_within_canvas': True,
          'pdf_text_overflow': False, 'source_links_present': True,
          'original_pdf_crop_images':4, 'pptx_embedded_images_match_crops':True,
          'speaker_notes_present': True, 'rendering': 'shared coordinates; independent PDF renderer',
          'native_powerpoint_render_verified': False}
(OUT / 'checks.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False))
